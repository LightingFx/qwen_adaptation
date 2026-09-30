"""Unblind the Tibetan-to-Chinese manual check and summarize it per system."""
import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHEET = ROOT / "manual_check/bo2zh_manual_check_evaluated.xlsx"
KEY = ROOT / "manual_check/bo2zh_key.csv"
OUT = ROOT / "results/ipm/bo2zh_manual_check_summary.json"
ORDER = ["qwen3_cpt_sft", "qwen3_va_sft_6epoch", "qwen2.5_cpt_sft"]


def wilson(k: int, n: int, z: float = 1.96) -> list[float]:
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return [round(100 * (centre - half), 1), round(100 * (centre + half), 1)]


key = {int(r["编号"]): r for r in csv.DictReader(KEY.open(encoding="utf-8-sig"))}
sheet = load_workbook(SHEET)["标注"]
rows = []
for idx, item, slot, _, _, _, lang, quality, _ in sheet.iter_rows(min_row=2, values_only=True):
    k = key[int(idx)]
    assert int(k["item_id"]) == int(item) and int(k["输出序号"]) == int(slot), idx
    rows.append({"system": k["system"], "lang": lang, "quality": str(quality),
                 "auto_han": int(k["auto_has_chinese"])})

inconsistent = [r for r in rows if (r["lang"] == "中文") != (r["quality"] in {"0", "1", "2"})]
summary = {"items_per_system": 100, "label_inconsistencies": len(inconsistent), "systems": {}}
by = defaultdict(list)
for r in rows:
    by[r["system"]].append(r)
for name in ORDER:
    rs = by[name]
    n = len(rs)
    lang = Counter(r["lang"] for r in rs)
    qual = Counter(r["quality"] for r in rs)
    chinese = lang["中文"]
    adequate = qual["2"]
    usable = qual["2"] + qual["1"]
    auto = sum(r["auto_han"] for r in rs)
    auto_pos_human_chinese = sum(1 for r in rs if r["auto_han"] and r["lang"] == "中文")
    summary["systems"][name] = {
        "n": n,
        "output_language": dict(lang),
        "quality": dict(qual),
        "chinese_translation_pct": round(100 * chinese / n, 1), "chinese_translation_ci": wilson(chinese, n),
        "usable_q1_or_q2_pct": round(100 * usable / n, 1), "usable_ci": wilson(usable, n),
        "adequate_q2_pct": round(100 * adequate / n, 1), "adequate_ci": wilson(adequate, n),
        "adequate_among_chinese_pct": round(100 * adequate / chinese, 1) if chinese else None,
        "auto_chinese_flag_pct": round(100 * auto / n, 1),
        "auto_flag_confirmed_as_chinese_translation": f"{auto_pos_human_chinese}/{auto}",
        "human_chinese_missed_by_auto_flag": sum(1 for r in rs if not r["auto_han"] and r["lang"] == "中文"),
    }
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
